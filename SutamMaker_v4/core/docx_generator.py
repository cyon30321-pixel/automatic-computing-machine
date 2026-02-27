"""
SutamMaker v4 — core/docx_generator.py
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
Word (DOCX) generation engine.
Depends on python-docx, and calls graph_engine for graph images.
No Tkinter dependencies.
"""

import os
import re
import shutil
import tempfile
import datetime

import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from config import CIRCLE_NUMS, build_paths
from core.graph_engine import process_and_draw_graph


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# PDF conversion helper
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def convert_docx_pairs_to_pdf(pairs):
    """Convert a list of (docx_path, pdf_target_path) pairs using docx2pdf."""
    try:
        from docx2pdf import convert as docx2pdf_convert
    except ImportError:
        return
    tmp_dir = tempfile.mkdtemp(prefix="exam_pdf_")
    mapping = []
    try:
        for idx, (dx, px) in enumerate(pairs, 1):
            if not os.path.exists(dx):
                continue
            tmp_dx = os.path.join(tmp_dir, f"file{idx}.docx")
            tmp_px = os.path.join(tmp_dir, f"file{idx}.pdf")
            shutil.copy2(dx, tmp_dx)
            mapping.append((tmp_px, px))
        docx2pdf_convert(tmp_dir)
        for tmp_px, target_px in mapping:
            if os.path.exists(tmp_px):
                os.makedirs(os.path.dirname(target_px), exist_ok=True)
                shutil.copy2(tmp_px, target_px)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# Internal helpers
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def _apply_doc_style(doc, font_name, font_size):
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(int(font_size))
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(1)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(1.2)
    sec.left_margin = sec.right_margin = Cm(1.3)


def _set_two_columns(section):
    try:
        spr = section._sectPr
        existing = spr.xpath("./w:cols")
        elem = existing[0] if existing else OxmlElement("w:cols")
        elem.set(qn("w:num"), "2")
        elem.set(qn("w:space"), "708")
        elem.set(qn("w:sep"), "1")
        if not existing:
            spr.append(elem)
    except Exception:
        pass


def _set_cell_bg(cell, color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def _add_watermark(doc, text):
    if not text:
        return
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(48)
        run.font.color.rgb = RGBColor(220, 220, 220)


def _add_header_footer(doc, header_text="", footer_text=""):
    for section in doc.sections:
        if header_text:
            h = section.header
            h.is_linked_to_previous = False
            hp = h.paragraphs[0] if h.paragraphs else h.add_paragraph()
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            hr = hp.add_run(header_text)
            hr.font.size = Pt(8)
            hr.font.color.rgb = RGBColor(150, 150, 150)
        if footer_text:
            f = section.footer
            f.is_linked_to_previous = False
            fp = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fr = fp.add_run(footer_text)
            fr.font.size = Pt(8)
            fr.font.color.rgb = RGBColor(150, 150, 150)


def _add_page_number_field(paragraph, font_size_pt=8, font_color_hex="969696"):
    """Insert a real page-number field via w:fldSimple."""
    prefix_run = paragraph.add_run("Page ")
    prefix_run.font.size = Pt(font_size_pt)
    prefix_run.font.color.rgb = RGBColor(
        int(font_color_hex[0:2], 16),
        int(font_color_hex[2:4], 16),
        int(font_color_hex[4:6], 16),
    )
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE \\* MERGEFORMAT ")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(font_size_pt * 2))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(font_size_pt * 2))
    rPr.append(szCs)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), font_color_hex)
    rPr.append(color_el)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._element.append(fld)


def _add_header_footer_v2(doc, hf_config, logo_path=""):
    """Enhanced header/footer — custom fields + OxmlElement page numbers.
    Applied to every section (including continuous ones).
    """
    academy = hf_config.get("academy_name", "")
    exam_name = hf_config.get("exam_name", "")
    contact = hf_config.get("contact_info", "")
    copyright_text = hf_config.get("copyright_text", "")
    show_page = hf_config.get("show_page_number", False)
    show_date = hf_config.get("show_date", False)
    date_str = datetime.datetime.now().strftime("%Y-%m-%d") if show_date else ""

    header_parts = [p for p in [academy, exam_name, date_str] if p]
    header_text = " | ".join(header_parts)
    footer_parts = [p for p in [copyright_text, contact] if p]
    footer_text = " | ".join(footer_parts)

    has_header = bool(header_text) or show_page
    has_footer = bool(footer_text) or show_page

    for section in doc.sections:
        if has_header:
            h = section.header
            h.is_linked_to_previous = False
            hp = h.paragraphs[0] if h.paragraphs else h.add_paragraph()
            hp.clear()
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if header_text:
                hr = hp.add_run(header_text)
                hr.font.size = Pt(8)
                hr.font.color.rgb = RGBColor(150, 150, 150)
            if show_page:
                if header_text:
                    sep = hp.add_run(" | ")
                    sep.font.size = Pt(8)
                    sep.font.color.rgb = RGBColor(150, 150, 150)
                _add_page_number_field(hp, 8, "969696")

        if has_footer:
            f = section.footer
            f.is_linked_to_previous = False
            fp = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
            fp.clear()
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if footer_text:
                fr = fp.add_run(footer_text)
                fr.font.size = Pt(8)
                fr.font.color.rgb = RGBColor(150, 150, 150)
            if show_page:
                if footer_text:
                    sep = fp.add_run(" | ")
                    sep.font.size = Pt(8)
                    sep.font.color.rgb = RGBColor(150, 150, 150)
                _add_page_number_field(fp, 8, "969696")


def _tight_para(doc, left_indent=None, space_before=0, space_after=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if left_indent:
        pf.left_indent = left_indent
    return p


def _tight_table_cell(cell, font_size=None):
    for p in cell.paragraphs:
        pf = p.paragraph_format
        pf.space_before = Pt(1)
        pf.space_after = Pt(1)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if font_size:
            for run in p.runs:
                run.font.size = Pt(font_size)


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# Public: create_exam_docx
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def create_exam_docx(target_dir, filename, exam_data, font_name, font_size,
                     logo_path, title, watermark="", header="", footer="",
                     numbering_start=1, hf_config=None):
    fs = int(font_size)
    doc = docx.Document()
    _apply_doc_style(doc, font_name, font_size)

    if watermark:
        _add_watermark(doc, watermark)
    if hf_config and any(hf_config.get(k) for k in (
        "academy_name", "exam_name", "show_page_number",
        "show_date", "copyright_text", "contact_info",
    )):
        _add_header_footer_v2(doc, hf_config, logo_path=logo_path)
    elif header or footer:
        _add_header_footer(doc, header, footer)

    if logo_path and os.path.exists(logo_path):
        lp = _tight_para(doc, space_after=2)
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.add_run().add_picture(logo_path, height=Cm(2.5))

    tp = doc.add_heading(title, 1)
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.space_after = Pt(4)

    sep = _tight_para(doc, space_before=0, space_after=6)
    sep_run = sep.add_run("━" * 60)
    sep_run.font.size = Pt(6)
    sep_run.font.color.rgb = RGBColor(180, 180, 180)
    sep.alignment = WD_ALIGN_PARAGRAPH.CENTER

    new_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    _set_two_columns(new_sec)

    for q_idx, q in enumerate(exam_data):
        display_num = numbering_start + q_idx
        parts = q["text"].split("[TABLE_PLACEHOLDER]")

        for idx, part in enumerate(parts):
            if part.strip() or idx == 0:
                p = _tight_para(doc, space_before=8 if idx == 0 else 1, space_after=2)
                if idx == 0:
                    num_run = p.add_run(f"{display_num}. ")
                    num_run.bold = True
                    num_run.font.size = Pt(fs + 1)
                text = part.strip()
                if text:
                    body_run = p.add_run(text)
                    body_run.font.size = Pt(fs)

            if idx < len(q["tables"]):
                td = q["tables"][idx]
                if td and td[0]:
                    tbl = doc.add_table(rows=len(td), cols=len(td[0]))
                    tbl.style = "Table Grid"
                    for ri, row in enumerate(td):
                        for ci, cell_text in enumerate(row):
                            if ci < len(tbl.columns):
                                cell = tbl.cell(ri, ci)
                                cell.text = cell_text
                                _tight_table_cell(cell, font_size=fs - 1)
                                if ri == 0:
                                    _set_cell_bg(cell, "E7E6E6")
                                    if cell.paragraphs[0].runs:
                                        cell.paragraphs[0].runs[0].bold = True
                                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        # jesi (passage box)
        if q["jesi"]:
            jt = doc.add_table(rows=1, cols=1)
            jt.style = "Table Grid"
            cell = jt.cell(0, 0)
            _set_cell_bg(cell, "F9F9F9")
            for ji, jline in enumerate(q["jesi"].split("\n")):
                if ji == 0:
                    cell.paragraphs[0].text = jline.strip()
                else:
                    cell.add_paragraph(jline.strip())
            _tight_table_cell(cell, font_size=fs - 1)

        # user local image
        if q.get("image_path") and os.path.exists(q["image_path"]):
            try:
                img_p = _tight_para(doc, space_before=2, space_after=2)
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.add_run().add_picture(q["image_path"], width=Cm(12.0))
            except Exception:
                pass

        # auto-generated graph
        if q.get("graph_tag"):
            tmp_img = os.path.join(target_dir, f"_tmp_graph_{q['num']}.png")
            img = process_and_draw_graph(q["graph_tag"], tmp_img)
            if img and os.path.exists(img):
                ip = _tight_para(doc, space_before=2, space_after=2)
                ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
                ip.add_run().add_picture(img, width=Cm(6.5))
                os.remove(img)

        # bogi box
        if q["bogi"]:
            bp = _tight_para(doc, space_before=3, space_after=1)
            bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            br = bp.add_run("< 보 기 >")
            br.bold = True
            br.font.size = Pt(fs)

            bt = doc.add_table(rows=1, cols=1)
            bt.style = "Table Grid"
            cell = bt.cell(0, 0)
            _set_cell_bg(cell, "FAFAFA")
            for bi, bline in enumerate(q["bogi"].split("\n")):
                if bi == 0:
                    cell.paragraphs[0].text = bline.strip()
                else:
                    cell.add_paragraph(bline.strip())
            _tight_table_cell(cell, font_size=fs - 1)

        # choices
        if q["choices"]:
            valid = [(i, c) for i, c in enumerate(q["choices"]) if c]
            if valid:
                max_len = max(len(c) for _, c in valid)
                if max_len <= 8 and len(valid) == 5:
                    cp = _tight_para(doc, left_indent=Cm(0.3), space_before=3, space_after=1)
                    parts_str = "   ".join(f"{CIRCLE_NUMS[i]} {c}" for i, c in valid)
                    cr = cp.add_run(parts_str)
                    cr.font.size = Pt(fs)
                elif max_len <= 18 and len(valid) >= 4:
                    first_row = valid[:3]
                    second_row = valid[3:]
                    cp1 = _tight_para(doc, left_indent=Cm(0.3), space_before=3, space_after=0)
                    r1 = cp1.add_run("   ".join(f"{CIRCLE_NUMS[i]} {c}" for i, c in first_row))
                    r1.font.size = Pt(fs)
                    if second_row:
                        cp2 = _tight_para(doc, left_indent=Cm(0.3), space_before=0, space_after=1)
                        r2 = cp2.add_run("   ".join(f"{CIRCLE_NUMS[i]} {c}" for i, c in second_row))
                        r2.font.size = Pt(fs)
                else:
                    for ci, (i, c) in enumerate(valid):
                        cp = _tight_para(doc, left_indent=Cm(0.5),
                                         space_before=1 if ci == 0 else 0, space_after=0)
                        cr = cp.add_run(f"{CIRCLE_NUMS[i]} {c}")
                        cr.font.size = Pt(fs)

        # divider between questions
        if q_idx < len(exam_data) - 1:
            div = _tight_para(doc, space_before=4, space_after=2)
            div_run = div.add_run("─" * 35)
            div_run.font.size = Pt(4)
            div_run.font.color.rgb = RGBColor(210, 210, 210)

    dx, px = build_paths(target_dir, filename)
    doc.save(dx)
    return dx, px


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# Public: create_answer_docx
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def create_answer_docx(target_dir, filename, raw_ans, exam_data, font_name, font_size,
                       logo_path, title, numbering_start=1, hf_config=None):
    doc = docx.Document()
    _apply_doc_style(doc, font_name, font_size)

    if hf_config and any(hf_config.get(k) for k in (
        "academy_name", "exam_name", "show_page_number",
        "show_date", "copyright_text", "contact_info",
    )):
        _add_header_footer_v2(doc, hf_config, logo_path=logo_path)

    if logo_path and os.path.exists(logo_path):
        lp = _tight_para(doc, space_after=2)
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.add_run().add_picture(logo_path, height=Cm(2.5))

    tp = doc.add_heading(f"정답 및 해설 ({title})", 1)
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.space_after = Pt(4)

    new_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    _set_two_columns(new_sec)

    ans_matches = re.findall(
        r"\b(\d+)(?:번|[\)\.])?(?:\s*정답\s*:?)?\s*[\(\[]?([①②③④⑤])[\)\]]?",
        raw_ans,
    )
    if ans_matches:
        p_title = _tight_para(doc, space_before=2, space_after=3)
        r = p_title.add_run("[ OMR 정답 카드 ]")
        r.bold = True
        r.font.size = Pt(12)
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        tbl = doc.add_table(rows=1, cols=10)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        for i in range(5):
            hdr[i * 2].text = "번호"
            hdr[i * 2 + 1].text = "정답"
            for c in (hdr[i * 2], hdr[i * 2 + 1]):
                _set_cell_bg(c, "2563EB")
                _tight_table_cell(c)
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if p.runs:
                        p.runs[0].bold = True
                        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

        row_cells = None
        for i, (qn_str, ans) in enumerate(ans_matches):
            if i % 5 == 0:
                row_cells = tbl.add_row().cells
            row_cells[(i % 5) * 2].text = qn_str
            row_cells[(i % 5) * 2 + 1].text = ans
            _set_cell_bg(row_cells[(i % 5) * 2], "EFF6FF")
            for p in (
                row_cells[(i % 5) * 2].paragraphs[0],
                row_cells[(i % 5) * 2 + 1].paragraphs[0],
            ):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        _tight_para(doc, space_before=4, space_after=4)

    for line in raw_ans.split("\n"):
        ls = line.strip()
        if ls:
            p = _tight_para(doc, space_before=1, space_after=1)
            p.add_run(ls)

    dx, px = build_paths(target_dir, filename, prefix="답안지_")
    doc.save(dx)
    return dx, px
