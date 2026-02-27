"""
통합 문제 은행 시스템 PRO v3.0
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
v2 대비 개선:
  - DB 컨텍스트 매니저 (연결 누수 방지)
  - 실제 DB 백업/복원 기능
  - 고급 필터 검색 (학교급·유형·년도 드롭다운)
  - 학생 삭제·이름 수정
  - 상태 표시줄 (DB 통계 실시간)
  - 종료 시 장바구니 확인
  - 키보드 단축키 (Ctrl+F 검색, Ctrl+S 생성)
  - 워크북 모드 버튼 추가
  - 정답지 지문별 개별 해설 입력 지원
"""

import os
import re
import sqlite3
import json
import datetime
import shutil
import platform
import subprocess
import time
import textwrap
import tkinter as tk
from tkinter import messagebox, filedialog, ttk
from contextlib import contextmanager
from collections import Counter
import csv

import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm

try:
    from docx2pdf import convert as _convert_pdf
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 1. 환경 설정
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(SCRIPT_DIR, "exam_pro_config.json")

DEFAULT_CONFIG = {
    "last_dir": SCRIPT_DIR,
    "db_dir": SCRIPT_DIR,
    "font_name": "맑은 고딕",
    "font_size": "10",
}


def _setup_korean_font():
    """matplotlib 한글 폰트를 OS에 맞게 자동 설정."""
    try:
        plt.style.use("seaborn-v0_8-whitegrid")
    except Exception:
        pass
    names = {f.name for f in fm.fontManager.ttflist}
    system = platform.system()
    if system == "Windows":
        for candidate in ("Malgun Gothic", "NanumGothic", "Gulim"):
            if candidate in names:
                plt.rcParams["font.family"] = candidate
                break
    elif system == "Darwin":
        plt.rcParams["font.family"] = "AppleGothic"
    else:
        plt.rcParams["font.family"] = "NanumGothic"
    plt.rcParams["axes.unicode_minus"] = False


_setup_korean_font()


def load_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                return {**DEFAULT_CONFIG, **json.load(f)}
        except (json.JSONDecodeError, OSError):
            pass
    return DEFAULT_CONFIG.copy()


def save_config(cfg):
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


def get_db_path(cfg):
    folder = cfg.get("db_dir", SCRIPT_DIR)
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, "question_bank_pro_v5.db")


def open_directory(path):
    if platform.system() == "Windows":
        os.startfile(path)
    elif platform.system() == "Darwin":
        subprocess.Popen(["open", path])
    else:
        subprocess.Popen(["xdg-open", path])


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 2. DB 레이어 — 컨텍스트 매니저 + 마이그레이션
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
@contextmanager
def db_conn(cfg):
    """예외 발생 시에도 안전하게 닫히는 DB 연결."""
    conn = sqlite3.connect(get_db_path(cfg))
    try:
        yield conn
    except Exception:
        conn.rollback()
        raise
    else:
        conn.commit()
    finally:
        conn.close()


def init_db(cfg):
    with db_conn(cfg) as conn:
        cur = conn.cursor()
        cur.execute("CREATE TABLE IF NOT EXISTS passages (id INTEGER PRIMARY KEY, content TEXT)")
        cur.execute("CREATE TABLE IF NOT EXISTS questions (id INTEGER PRIMARY KEY, passage_id INTEGER)")
        cur.execute("CREATE TABLE IF NOT EXISTS students (id INTEGER PRIMARY KEY, name TEXT)")
        cur.execute("CREATE TABLE IF NOT EXISTS analysis_records (id INTEGER PRIMARY KEY, student_id INTEGER)")

        def ensure(table, col, ctype):
            cur.execute(f"PRAGMA table_info({table})")
            existing = {row[1] for row in cur.fetchall()}
            if col not in existing:
                try:
                    cur.execute(f"ALTER TABLE {table} ADD COLUMN {col} {ctype}")
                except sqlite3.OperationalError:
                    pass

        for c, t in [
            ("category1", "TEXT"), ("category2", "TEXT"), ("school_year", "TEXT"),
            ("exam_year", "TEXT"), ("exam_month", "TEXT"), ("publisher", "TEXT"),
            ("extra_tags", "TEXT"), ("answer_text", "TEXT"), ("created_at", "TEXT"),
        ]:
            ensure("passages", c, t)

        for c, t in [("q_num", "TEXT"), ("content", "TEXT"), ("usage_count", "INTEGER DEFAULT 0")]:
            ensure("questions", c, t)

        ensure("students", "grade", "TEXT")

        for c, t in [("record_date", "TEXT"), ("weakness_tag", "TEXT"), ("score", "INTEGER"), ("feedback", "TEXT")]:
            ensure("analysis_records", c, t)


def db_stats(cfg):
    """DB 요약 통계를 딕셔너리로 반환."""
    with db_conn(cfg) as conn:
        cur = conn.cursor()
        p = cur.execute("SELECT COUNT(*) FROM passages").fetchone()[0]
        q = cur.execute("SELECT COUNT(*) FROM questions").fetchone()[0]
        s = cur.execute("SELECT COUNT(*) FROM students").fetchone()[0]
        return {"passages": p, "questions": q, "students": s}


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 3. PDF 변환
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def convert_to_pdf_safe(docx_path, pdf_path):
    time.sleep(1.0)
    if platform.system() == "Windows":
        try:
            import win32com.client
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(docx_path, ReadOnly=True)
            doc.SaveAs(pdf_path, FileFormat=17)
            doc.Close()
            word.Quit()
            return True, "성공"
        except Exception as e:
            return False, f"Word COM 오류: {e}"
    else:
        if not PDF_AVAILABLE:
            return False, "docx2pdf 미설치"
        try:
            _convert_pdf(docx_path, pdf_path)
            return True, "성공"
        except Exception as e:
            return False, str(e)


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 4. Word 문서 생성 엔진
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def _setup_doc(doc, cfg, title):
    style = doc.styles["Normal"]
    style.font.name = cfg["font_name"]
    style._element.rPr.rFonts.set(qn("w:eastAsia"), cfg["font_name"])
    style.font.size = Pt(int(cfg["font_size"]))
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    sec.left_margin = sec.right_margin = Cm(1.5)
    doc.add_heading(title, 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def create_files(target_dir, prefix, exam_data, cfg, is_workbook=False, output_format="Word + PDF"):
    ts = datetime.datetime.now().strftime("%m%d_%H%M")

    # ── 문제지 ────────────────────────────────────
    doc_q = docx.Document()
    _setup_doc(doc_q, cfg, f"실전 대비 맞춤형 {'워크북' if is_workbook else '시험지'}")

    if not is_workbook:
        new_sec = doc_q.add_section(WD_SECTION.CONTINUOUS)
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "708")
        new_sec._sectPr.append(cols)

    g_idx = 1
    for idx, item in enumerate(exam_data, 1):
        if not item["passage"]:
            continue
        item["q_mappings"] = []

        if is_workbook:
            h = doc_q.add_paragraph()
            h.add_run(f"■ 지문 {idx}  ").bold = True
            h.add_run(f"[{item.get('info', '')}]").font.color.rgb = RGBColor(100, 100, 100)

            sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", item["passage"].strip()) if s.strip()]
            for s_i, sent in enumerate(sentences, 1):
                p = doc_q.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.add_run(f"{s_i}. {sent}").bold = True
                for _ in range(2):
                    bl = doc_q.add_paragraph()
                    bl.paragraph_format.space_after = Pt(0)
                    bl.add_run("_" * 74).font.color.rgb = RGBColor(220, 220, 220)
            doc_q.add_paragraph("\n")
        else:
            intro = doc_q.add_paragraph()
            intro.add_run("※ 다음 글을 읽고 물음에 답하시오.").bold = True
            tbl = doc_q.add_table(rows=1, cols=1)
            tbl.style = "Table Grid"
            tbl.cell(0, 0).text = item["passage"]
            doc_q.add_paragraph()

            for q in item["questions"]:
                p = doc_q.add_paragraph()
                p.add_run(f"{g_idx}. ").bold = True
                p.add_run(q["text"])
                doc_q.add_paragraph()
                item["q_mappings"].append({
                    "new_num": g_idx,
                    "orig_num": q["num"],
                    "preview": q["text"][:25].replace("\n", " ") + "...",
                })
                g_idx += 1

    q_docx = os.path.normpath(os.path.join(target_dir, f"{prefix}_문제지_{ts}.docx"))
    doc_q.save(q_docx)

    # ── 정답지 ────────────────────────────────────
    doc_a = docx.Document()
    _setup_doc(doc_a, cfg, f"맞춤형 {'워크북' if is_workbook else '시험지'} — 정답 및 해설")

    for idx, item in enumerate(exam_data, 1):
        h = doc_a.add_paragraph()
        h.add_run(f"■ [지문 {idx}] 해설\n").bold = True

        if not is_workbook and item.get("q_mappings"):
            pm = doc_a.add_paragraph()
            rt = pm.add_run("[※ 문항 번호 매칭표]\n")
            rt.bold = True
            rt.font.color.rgb = RGBColor(0, 112, 192)

            lines = ""
            for m in item["q_mappings"]:
                orig = f"{m['orig_num']}번" if m["orig_num"] != "-" else "서술형"
                lines += f"▶ 시험지 {m['new_num']}번 = 본문 해설 {orig} 참조 | {m['preview']}\n"
            rm = pm.add_run(lines)
            rm.font.color.rgb = RGBColor(100, 100, 100)
            rm.font.size = Pt(9)
            doc_a.add_paragraph()

        body = doc_a.add_paragraph()
        ans = (item.get("answer_text") or "").strip()
        body.add_run(ans if ans else "(입력된 정답/해설이 없습니다.)")
        doc_a.add_paragraph("\n" + "=" * 50 + "\n")

    a_docx = os.path.normpath(os.path.join(target_dir, f"{prefix}_정답지_{ts}.docx"))
    doc_a.save(a_docx)

    # ── PDF 변환 ──────────────────────────────────
    msg = "Word 파일(문제/답안) 분리 생성 완료!"
    if output_format in ("PDF만", "Word + PDF"):
        q_pdf = q_docx.replace(".docx", ".pdf")
        a_pdf = a_docx.replace(".docx", ".pdf")
        q_ok, q_err = convert_to_pdf_safe(q_docx, q_pdf)
        a_ok, a_err = convert_to_pdf_safe(a_docx, a_pdf)

        if q_ok and a_ok:
            msg = "Word + PDF 분리 생성 완료!"
            if output_format == "PDF만":
                for f in (q_docx, a_docx):
                    try:
                        os.remove(f)
                    except OSError:
                        pass
                msg = "PDF 분리 생성 완료!"
        else:
            msg = f"Word는 생성했으나 PDF 변환 실패: {q_err}"
    return msg


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 5. 메인 앱
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
class ExamProApp:
    def __init__(self, root):
        self.root = root
        self.root.title("통합 문제 은행 시스템 PRO v3")
        self.root.geometry("1320x960")
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

        self.cfg = load_config()
        init_db(self.cfg)

        self.cart = {}
        self.current_student_id = None
        self.current_student_name = ""
        self.analyzed_weakness = None
        self.tooltip_win = None
        self.preview_data = {}

        self._build_ui()
        self._bind_shortcuts()
        self._refresh_status()

    # ──────────────────────────────────────────────
    #  UI 프레임워크
    # ──────────────────────────────────────────────
    def _build_ui(self):
        # 상태바 (하단)
        self.status_bar = tk.Frame(self.root, bg="#e2e8f0", height=26)
        self.status_bar.pack(fill="x", side="bottom")
        self.status_bar.pack_propagate(False)
        self.lbl_status = tk.Label(
            self.status_bar, text="", font=("맑은 고딕", 8),
            bg="#e2e8f0", fg="#475569",
        )
        self.lbl_status.pack(side="left", padx=12)

        # 탭
        self.tabs = ttk.Notebook(self.root)
        self.tabs.pack(fill="both", expand=True, padx=8, pady=8)

        self.tab_input = ttk.Frame(self.tabs)
        self.tab_bank = ttk.Frame(self.tabs)
        self.tab_analysis = ttk.Frame(self.tabs)
        self.tab_config = ttk.Frame(self.tabs)

        self.tabs.add(self.tab_input, text="  데이터 입력  ")
        self.tabs.add(self.tab_bank, text="  문제 관리 및 출력  ")
        self.tabs.add(self.tab_analysis, text="  AI 오답 분석  ")
        self.tabs.add(self.tab_config, text="  설정 및 백업  ")

        self._build_input_tab()
        self._build_bank_tab()
        self._build_analysis_tab()
        self._build_config_tab()

    def _bind_shortcuts(self):
        self.root.bind("<Control-f>", lambda e: (self.tabs.select(self.tab_bank), self.ent_search.focus_set()))
        self.root.bind("<Control-s>", lambda e: self._generate("exam"))

    def _set_status(self, text):
        self.lbl_status.config(text=text)

    def _refresh_status(self):
        try:
            s = db_stats(self.cfg)
            cart_count = sum(1 if not v else len(v) for v in self.cart.values())
            self._set_status(
                f"DB: 지문 {s['passages']}개 | 문제 {s['questions']}개 | "
                f"학생 {s['students']}명    장바구니: {cart_count}건"
            )
        except Exception:
            self._set_status("DB 연결 확인 필요")

    def _on_close(self):
        if self.cart:
            if not messagebox.askyesno("종료 확인", "장바구니에 항목이 있습니다.\n정말 종료하시겠습니까?"):
                return
        self.root.destroy()

    # ──────────────────────────────────────────────
    #  TAB 1: 데이터 입력
    # ──────────────────────────────────────────────
    def _build_input_tab(self):
        main = tk.Frame(self.tab_input)
        main.pack(fill="both", expand=True, padx=20, pady=15)

        # 분류 태그
        tag_frame = tk.LabelFrame(main, text=" 분류 태그 ", font=("맑은 고딕", 10, "bold"), pady=8)
        tag_frame.pack(fill="x", pady=(0, 8))

        r1 = tk.Frame(tag_frame)
        r1.pack(fill="x", padx=10)

        def _combo(parent, label, values, width=7):
            tk.Label(parent, text=label).pack(side="left", padx=(8, 0))
            cb = ttk.Combobox(parent, values=values, width=width, state="readonly")
            cb.pack(side="left", padx=4)
            return cb

        self.combo_cat1 = _combo(r1, "학교급:", ["중학교", "고등학교"], 6)
        self.combo_year = _combo(r1, "학년:", ["1학년", "2학년", "3학년", "공통"], 6)
        self.combo_cat2 = _combo(r1, "유형:", ["내신", "모의고사", "부교재"], 8)
        self.combo_exam_year = _combo(r1, "년도:", ["2026년", "2025년", "2024년", "2023년"], 7)
        self.combo_exam_month = _combo(r1, "월:", ["3월", "4월", "6월", "9월", "11월"], 5)

        r2 = tk.Frame(tag_frame)
        r2.pack(fill="x", padx=10, pady=8)
        tk.Label(r2, text="학교명/출판사:").pack(side="left")
        self.ent_pub = tk.Entry(r2, width=20)
        self.ent_pub.pack(side="left", padx=5)
        tk.Label(r2, text="추가 태그:").pack(side="left", padx=(10, 0))
        self.ent_extra_tags = tk.Entry(r2, width=30)
        self.ent_extra_tags.pack(side="left", padx=5)

        # 지문 입력
        tk.Label(
            main, text="지문+문제 붙여넣기  (지문과 문제 사이 빈 줄 필수)",
            font=("맑은 고딕", 10, "bold"), fg="#2563eb",
        ).pack(anchor="w", pady=(5, 0))
        self.txt_input = tk.Text(main, font=("Consolas", 10), height=14)
        self.txt_input.pack(fill="both", expand=True, pady=4)

        # 정답
        tk.Label(
            main, text="정답 및 해설 (선택 — 별도 정답지로 출력)",
            font=("맑은 고딕", 10, "bold"), fg="#16a34a",
        ).pack(anchor="w", pady=(4, 0))
        self.txt_answer = tk.Text(main, font=("Consolas", 10), height=5)
        self.txt_answer.pack(fill="x", pady=4)

        tk.Button(
            main, text="지문 자동 분해 및 검수", bg="#eab308", fg="#1e293b",
            font=("맑은 고딕", 12, "bold"), height=2,
            command=self._open_preview,
        ).pack(fill="x")

    def _open_preview(self):
        raw = self.txt_input.get("1.0", tk.END).strip()
        ans_raw = self.txt_answer.get("1.0", tk.END).strip()
        if not raw:
            return messagebox.showwarning("경고", "텍스트를 입력하세요.")

        blocks = re.split(r"\n(?=\s*\[\d+\]|\s*다음\s*글을\s*읽고)", "\n" + raw)
        parsed = []
        for block in blocks:
            block = block.strip()
            if len(block) < 15:
                continue
            parts = re.split(r"\n\s*\n(?=\d+\.\s)", block)
            passage = re.sub(
                r"^(?:\[\d+\]\s*)?(?:다음\s*글을\s*읽고\s*물음에\s*답하시오\.?)?",
                "", parts[0].strip(),
            ).strip()
            tag_m = re.match(r"^(\[\d+\])\s*", parts[0].strip())
            tag = tag_m.group(1) if tag_m else ""
            parsed.append({
                "passage": passage,
                "questions": parts[1:] if len(parts) > 1 else [],
                "tag": tag,
            })

        if not parsed:
            return messagebox.showwarning("오류", "유효한 지문을 찾지 못했습니다.")

        top = tk.Toplevel(self.root)
        top.title(f"검수 — {len(parsed)}개 지문 발견")
        top.geometry("920x720")
        top.attributes("-topmost", True)

        txt = tk.Text(top, font=("Consolas", 10), bg="#f8fafc")
        txt.pack(fill="both", expand=True, padx=16, pady=8)

        for i, item in enumerate(parsed, 1):
            txt.insert(tk.END, f"{'='*20} [지문 {i}] {'='*20}\n\n")
            txt.insert(tk.END, item["passage"] + "\n\n")
            for q in item["questions"]:
                txt.insert(tk.END, q.strip() + "\n\n")
            txt.insert(tk.END, "\n")

        def do_save():
            with db_conn(self.cfg) as conn:
                cur = conn.cursor()
                for item in parsed:
                    extra = f"{item['tag']} {self.ent_extra_tags.get().strip()}".strip() if item["tag"] else self.ent_extra_tags.get().strip()
                    cur.execute(
                        """INSERT INTO passages
                           (content,category1,category2,school_year,exam_year,
                            exam_month,publisher,extra_tags,answer_text,created_at)
                           VALUES (?,?,?,?,?,?,?,?,?,?)""",
                        (item["passage"], self.combo_cat1.get(), self.combo_cat2.get(),
                         self.combo_year.get(), self.combo_exam_year.get(),
                         self.combo_exam_month.get(), self.ent_pub.get(),
                         extra, ans_raw, datetime.date.today().isoformat()),
                    )
                    pid = cur.lastrowid
                    for qt in item["questions"]:
                        m = re.match(r"^(\d+)\.\s*(.*)", qt.strip(), re.DOTALL)
                        cur.execute(
                            "INSERT INTO questions (passage_id, q_num, content) VALUES (?,?,?)",
                            (pid, m.group(1) if m else "-", m.group(2) if m else qt.strip()),
                        )
            messagebox.showinfo("완료", f"{len(parsed)}개 지문 저장 완료!")
            top.destroy()
            self.txt_input.delete("1.0", tk.END)
            self.txt_answer.delete("1.0", tk.END)
            self._refresh_bank()
            self._refresh_status()

        tk.Button(
            top, text="최종 저장", bg="#16a34a", fg="white",
            font=("맑은 고딕", 11, "bold"), height=2, command=do_save,
        ).pack(fill="x", padx=16, pady=8)

    # ──────────────────────────────────────────────
    #  TAB 2: 문제 관리 & 장바구니
    # ──────────────────────────────────────────────
    def _build_bank_tab(self):
        # ── 검색 바 (텍스트 + 필터 드롭다운) ────
        search_frame = tk.Frame(self.tab_bank)
        search_frame.pack(fill="x", padx=10, pady=6)

        tk.Label(search_frame, text="검색:").pack(side="left")
        self.ent_search = tk.Entry(search_frame, width=22)
        self.ent_search.pack(side="left", padx=4)
        self.ent_search.bind("<Return>", lambda e: self._refresh_bank())

        tk.Label(search_frame, text="학교급:").pack(side="left", padx=(10, 0))
        self.f_cat1 = ttk.Combobox(search_frame, values=["전체", "중학교", "고등학교"], width=6, state="readonly")
        self.f_cat1.current(0)
        self.f_cat1.pack(side="left", padx=2)

        tk.Label(search_frame, text="유형:").pack(side="left", padx=(6, 0))
        self.f_cat2 = ttk.Combobox(search_frame, values=["전체", "내신", "모의고사", "부교재"], width=7, state="readonly")
        self.f_cat2.current(0)
        self.f_cat2.pack(side="left", padx=2)

        tk.Label(search_frame, text="년도:").pack(side="left", padx=(6, 0))
        self.f_year = ttk.Combobox(search_frame, values=["전체", "2026년", "2025년", "2024년", "2023년"], width=7, state="readonly")
        self.f_year.current(0)
        self.f_year.pack(side="left", padx=2)

        tk.Button(search_frame, text="조회", command=self._refresh_bank, bg="#6b7280", fg="white").pack(side="left", padx=8)
        tk.Button(search_frame, text="초기화", command=self._reset_filters).pack(side="left")

        # ── 트리뷰 ──────────────────────────
        tree_cols = ("ID", "구분", "분류/시기", "출판사/태그", "내용 요약", "등록일")
        self.tree = ttk.Treeview(
            self.tab_bank, columns=tree_cols,
            show="headings", height=10, selectmode="extended",
        )
        for col, w in zip(tree_cols, [60, 55, 140, 140, 420, 85]):
            self.tree.heading(col, text=col)
            self.tree.column(col, width=w)
        self.tree.pack(fill="x", padx=10, pady=4)

        self.tree.bind("<Motion>", self._show_tooltip)
        self.tree.bind("<Leave>", self._hide_tooltip)
        self.tree.bind("<Double-1>", lambda e: self._cart_add())

        # ── 트리 액션 버튼 ──────────────────
        tb = tk.Frame(self.tab_bank)
        tb.pack(fill="x", padx=10, pady=2)
        tk.Button(tb, text="선택 수정", command=self._edit_selected).pack(side="left", padx=2)
        tk.Button(tb, text="DB에서 삭제", bg="#dc2626", fg="white", font=("맑은 고딕", 9, "bold"), command=self._delete_selected).pack(side="left", padx=10)
        tk.Button(tb, text="선택 항목 장바구니 담기", bg="#2563eb", fg="white", font=("맑은 고딕", 9, "bold"), command=self._cart_add).pack(side="right", padx=2)

        # ── 장바구니 ────────────────────────
        cart_frame = tk.LabelFrame(self.tab_bank, text=" 장바구니 (더블클릭=삭제) ", font=("맑은 고딕", 9, "bold"))
        cart_frame.pack(fill="both", expand=True, padx=10, pady=8)

        self.list_cart = tk.Listbox(cart_frame, height=6, selectmode=tk.EXTENDED)
        self.list_cart.pack(side="left", fill="both", expand=True, padx=4, pady=4)
        self.list_cart.bind("<Double-1>", self._cart_remove)

        btn_col = tk.Frame(cart_frame)
        btn_col.pack(side="right", padx=8, fill="y")

        fmt = tk.Frame(btn_col)
        fmt.pack(fill="x", pady=4)
        tk.Label(fmt, text="포맷:").pack(side="left")
        self.combo_format = ttk.Combobox(fmt, values=["Word + PDF", "Word만", "PDF만"], width=12, state="readonly")
        self.combo_format.current(0)
        self.combo_format.pack(side="left", padx=4)

        tk.Button(
            btn_col, text="시험지 생성", bg="#16a34a", fg="white",
            font=("맑은 고딕", 10, "bold"), width=20,
            command=lambda: self._generate("exam"),
        ).pack(pady=3)
        tk.Button(
            btn_col, text="워크북 생성", bg="#7c3aed", fg="white",
            font=("맑은 고딕", 10, "bold"), width=20,
            command=lambda: self._generate("workbook"),
        ).pack(pady=3)

        act = tk.Frame(btn_col)
        act.pack(fill="x", pady=4)
        tk.Button(act, text="선택 빼기", command=self._cart_remove).pack(side="left", expand=True, fill="x", padx=1)
        tk.Button(act, text="전체 비우기", command=self._cart_clear).pack(side="left", expand=True, fill="x", padx=1)

    def _reset_filters(self):
        self.ent_search.delete(0, tk.END)
        self.f_cat1.current(0)
        self.f_cat2.current(0)
        self.f_year.current(0)
        self._refresh_bank()

    def _refresh_bank(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        self.preview_data.clear()

        search = self.ent_search.get().strip()
        fc1 = self.f_cat1.get()
        fc2 = self.f_cat2.get()
        fy = self.f_year.get()

        with db_conn(self.cfg) as conn:
            cur = conn.cursor()

            query = "SELECT id,category1,category2,school_year,exam_year,exam_month,publisher,extra_tags,content,created_at FROM passages WHERE 1=1"
            params = []

            if search:
                query += " AND (content LIKE ? OR extra_tags LIKE ? OR publisher LIKE ?)"
                params += [f"%{search}%", f"%{search}%", f"%{search}%"]
            if fc1 != "전체":
                query += " AND category1=?"
                params.append(fc1)
            if fc2 != "전체":
                query += " AND category2=?"
                params.append(fc2)
            if fy != "전체":
                query += " AND exam_year=?"
                params.append(fy)

            query += " ORDER BY id DESC"

            for row in cur.execute(query, params).fetchall():
                pid = row[0]
                iid_p = f"P_{pid}"
                period = f"{row[3] or ''} {row[4] or ''} {row[5] or ''}".strip()
                tags = f"{row[6] or ''} {row[7] or ''}".strip()
                summary = (row[8] or "")[:80].replace("\n", " ")

                p_node = self.tree.insert(
                    "", "end", iid=iid_p,
                    values=(f"P-{pid}", "지문", period, tags, summary, row[9] or ""),
                )
                self.preview_data[iid_p] = row[8]

                for qr in cur.execute("SELECT id,q_num,content FROM questions WHERE passage_id=?", (pid,)):
                    iid_q = f"Q_{qr[0]}_{pid}"
                    q_summary = f"[{qr[1]}번] {(qr[2] or '')[:80].replace(chr(10), ' ')}"
                    self.tree.insert(
                        p_node, "end", iid=iid_q,
                        values=(f"Q-{qr[0]}", "문제", "", "", q_summary, ""),
                    )
                    self.preview_data[iid_q] = qr[2]

    def _show_tooltip(self, event):
        item_id = self.tree.identify_row(event.y)
        col = self.tree.identify_column(event.x)
        if item_id and col == "#5":
            text = self.preview_data.get(item_id, "")
            if not text:
                return
            wrapped = textwrap.fill(text[:500] + ("..." if len(text) > 500 else ""), width=60)
            if self.tooltip_win:
                self._tip_label.config(text=wrapped)
                self.tooltip_win.geometry(f"+{event.x_root+15}+{event.y_root+15}")
            else:
                self.tooltip_win = tk.Toplevel(self.root)
                self.tooltip_win.wm_overrideredirect(True)
                self.tooltip_win.geometry(f"+{event.x_root+15}+{event.y_root+15}")
                self._tip_label = tk.Label(
                    self.tooltip_win, text=wrapped, justify="left",
                    bg="#fefce8", relief="solid", borderwidth=1, font=("맑은 고딕", 9),
                )
                self._tip_label.pack(ipadx=5, ipady=5)
        else:
            self._hide_tooltip()

    def _hide_tooltip(self, event=None):
        if self.tooltip_win:
            self.tooltip_win.destroy()
            self.tooltip_win = None

    def _delete_selected(self):
        sel = self.tree.selection()
        if not sel:
            return messagebox.showwarning("알림", "삭제할 항목을 선택하세요.")
        if not messagebox.askyesno("경고", f"{len(sel)}개 항목을 영구 삭제하시겠습니까?\n(지문 삭제 시 딸린 문제도 삭제)"):
            return
        with db_conn(self.cfg) as conn:
            cur = conn.cursor()
            for iid in sel:
                if iid.startswith("P_"):
                    pid = int(iid.split("_")[1])
                    cur.execute("DELETE FROM passages WHERE id=?", (pid,))
                    cur.execute("DELETE FROM questions WHERE passage_id=?", (pid,))
                    self.cart.pop(pid, None)
                elif iid.startswith("Q_"):
                    parts = iid.split("_")
                    qid, pid = int(parts[1]), int(parts[2])
                    cur.execute("DELETE FROM questions WHERE id=?", (qid,))
                    if pid in self.cart and self.cart[pid]:
                        self.cart[pid].discard(qid)
                        if not self.cart[pid]:
                            del self.cart[pid]
        self._refresh_bank()
        self._sync_cart_ui()
        self._refresh_status()
        messagebox.showinfo("완료", "삭제 완료.")

    def _edit_selected(self):
        sel = self.tree.selection()
        if not sel:
            return
        iid = sel[0]
        if iid.startswith("Q_"):
            return messagebox.showinfo("안내", "지문 단위로 선택 후 수정하세요.")
        pid = int(iid.split("_")[1])

        with db_conn(self.cfg) as conn:
            cur = conn.cursor()
            p_row = cur.execute("SELECT content, answer_text FROM passages WHERE id=?", (pid,)).fetchone()
            q_rows = cur.execute("SELECT q_num, content FROM questions WHERE passage_id=?", (pid,)).fetchall()

        if not p_row:
            return

        top = tk.Toplevel(self.root)
        top.title(f"수정 — 지문 P-{pid}")
        top.geometry("900x850")
        top.attributes("-topmost", True)

        tk.Label(top, text="지문", font=("맑은 고딕", 10, "bold")).pack(pady=4)
        txt_p = tk.Text(top, height=10, font=("Consolas", 10))
        txt_p.pack(fill="x", padx=16)
        txt_p.insert("1.0", p_row[0])

        tk.Label(top, text="문제 (사이에 빈 줄 필수)", font=("맑은 고딕", 10, "bold")).pack(pady=4)
        txt_q = tk.Text(top, height=10, font=("Consolas", 10))
        txt_q.pack(fill="x", padx=16)
        for q in q_rows:
            num = f"{q[0]}. " if q[0] != "-" else ""
            txt_q.insert(tk.END, f"{num}{q[1]}\n\n")

        tk.Label(top, text="정답 및 해설", font=("맑은 고딕", 10, "bold"), fg="#16a34a").pack(pady=4)
        txt_a = tk.Text(top, height=5, font=("Consolas", 10))
        txt_a.pack(fill="x", padx=16)
        txt_a.insert("1.0", p_row[1] or "")

        def do_save():
            with db_conn(self.cfg) as conn:
                cur = conn.cursor()
                cur.execute("UPDATE passages SET content=?, answer_text=? WHERE id=?",
                            (txt_p.get("1.0", tk.END).strip(), txt_a.get("1.0", tk.END).strip(), pid))
                cur.execute("DELETE FROM questions WHERE passage_id=?", (pid,))
                for qt in re.split(r"\n\s*\n(?=\d+\.\s)", txt_q.get("1.0", tk.END).strip()):
                    if not qt.strip():
                        continue
                    m = re.match(r"^(\d+)\.\s*(.*)", qt.strip(), re.DOTALL)
                    cur.execute(
                        "INSERT INTO questions (passage_id,q_num,content) VALUES (?,?,?)",
                        (pid, m.group(1) if m else "-", m.group(2) if m else qt.strip()),
                    )
            messagebox.showinfo("완료", "수정 반영 완료.")
            top.destroy()
            self._refresh_bank()

        tk.Button(
            top, text="수정 저장", bg="#2563eb", fg="white",
            font=("맑은 고딕", 11, "bold"), height=2, command=do_save,
        ).pack(fill="x", padx=16, pady=12)

    # ── 장바구니 조작 ─────────────────────────
    def _sync_cart_ui(self):
        self.list_cart.delete(0, tk.END)
        with db_conn(self.cfg) as conn:
            cur = conn.cursor()
            for pid, qids in self.cart.items():
                if not qids:
                    cnt = cur.execute("SELECT COUNT(*) FROM questions WHERE passage_id=?", (pid,)).fetchone()[0]
                    self.list_cart.insert("end", f"지문 전체 [ID:{pid}] ({cnt}문제)")
                else:
                    for qid in qids:
                        qn_row = cur.execute("SELECT q_num FROM questions WHERE id=?", (qid,)).fetchone()
                        num = qn_row[0] if qn_row else "-"
                        self.list_cart.insert("end", f"개별 문제 [지문{pid} - {num}번 (ID:{qid})]")
        self._refresh_status()

    def _cart_add(self, event=None):
        sel = self.tree.selection()
        if not sel:
            return messagebox.showwarning("알림", "항목을 선택하세요.")
        for iid in sel:
            if iid.startswith("P_"):
                pid = int(iid.split("_")[1])
                self.cart[pid] = set()
            elif iid.startswith("Q_"):
                parts = iid.split("_")
                qid, pid = int(parts[1]), int(parts[2])
                if pid not in self.cart:
                    self.cart[pid] = {qid}
                elif self.cart[pid]:
                    self.cart[pid].add(qid)
        self._sync_cart_ui()

    def _cart_remove(self, event=None):
        indices = list(self.list_cart.curselection())
        if not indices:
            return
        for idx in reversed(indices):
            text = self.list_cart.get(idx)
            m_full = re.search(r"\[ID:(\d+)\]", text)
            m_indiv = re.search(r"\(ID:(\d+)\)", text)
            if "지문 전체" in text and m_full:
                self.cart.pop(int(m_full.group(1)), None)
            elif m_indiv:
                qid = int(m_indiv.group(1))
                for pk, qs in list(self.cart.items()):
                    if qs and qid in qs:
                        qs.discard(qid)
                        if not qs:
                            del self.cart[pk]
        self._sync_cart_ui()

    def _cart_clear(self):
        self.cart.clear()
        self._sync_cart_ui()

    def _generate(self, mode):
        if not self.cart:
            return messagebox.showwarning("경고", "장바구니가 비어있습니다.")

        with db_conn(self.cfg) as conn:
            cur = conn.cursor()
            exam_data = []
            for pid, qids in self.cart.items():
                p = cur.execute(
                    "SELECT content,category2,school_year,exam_year,exam_month,extra_tags,answer_text FROM passages WHERE id=?",
                    (pid,),
                ).fetchone()
                if not p:
                    continue
                if qids:
                    ph = ",".join("?" * len(qids))
                    qs = cur.execute(f"SELECT q_num,content FROM questions WHERE passage_id=? AND id IN ({ph})", (pid, *qids)).fetchall()
                else:
                    qs = cur.execute("SELECT q_num,content FROM questions WHERE passage_id=?", (pid,)).fetchall()

                exam_data.append({
                    "pid": pid,
                    "passage": p[0],
                    "questions": [{"num": q[0], "text": q[1]} for q in qs],
                    "info": f"{p[2] or ''} {p[3] or ''} {p[4] or ''}".strip(),
                    "answer_text": p[6],
                })

        is_wb = mode != "exam"
        prefix = "Workbook" if is_wb else "Test"

        messagebox.showinfo("생성 시작", "파일 생성을 시작합니다.\nPDF 변환 시 2~3초 소요될 수 있습니다.")
        result = create_files(self.cfg["last_dir"], prefix, exam_data, self.cfg, is_workbook=is_wb, output_format=self.combo_format.get())

        messagebox.showinfo("완료", f"{result}\n\n저장: {self.cfg['last_dir']}")
        self._cart_clear()
        open_directory(self.cfg["last_dir"])

    # ──────────────────────────────────────────────
    #  TAB 3: AI 오답 분석
    # ──────────────────────────────────────────────
    def _build_analysis_tab(self):
        self.tab_analysis.columnconfigure(0, weight=1)
        self.tab_analysis.columnconfigure(1, weight=3)
        self.tab_analysis.columnconfigure(2, weight=2)

        # 좌: 학생 선택
        f_stu = tk.LabelFrame(self.tab_analysis, text=" 1. 학생 선택 ", font=("맑은 고딕", 10, "bold"))
        f_stu.grid(row=0, column=0, sticky="nsew", padx=8, pady=8)

        self.ent_student = tk.Entry(f_stu)
        self.ent_student.pack(fill="x", padx=6, pady=4)
        self.ent_student.bind("<Return>", lambda e: self._add_student())

        btn_stu = tk.Frame(f_stu)
        btn_stu.pack(fill="x", padx=6, pady=2)
        tk.Button(btn_stu, text="추가", command=self._add_student, width=8).pack(side="left", padx=2)
        tk.Button(btn_stu, text="삭제", command=self._delete_student, bg="#dc2626", fg="white", width=8).pack(side="left", padx=2)
        tk.Button(btn_stu, text="이름 수정", command=self._rename_student, width=8).pack(side="left", padx=2)

        self.list_students = tk.Listbox(f_stu)
        self.list_students.pack(fill="both", expand=True, padx=6, pady=6)
        self.list_students.bind("<<ListboxSelect>>", self._on_student_select)
        self._refresh_students()

        tk.Button(
            f_stu, text="AI 프롬프트 복사", bg="#6b7280", fg="white",
            font=("맑은 고딕", 9, "bold"), command=self._copy_ai_prompt,
        ).pack(fill="x", padx=6, pady=4)

        # 중: AI 분석 입력
        f_ai = tk.LabelFrame(self.tab_analysis, text=" 2. AI 분석 결과 입력 ", font=("맑은 고딕", 10, "bold"))
        f_ai.grid(row=0, column=1, sticky="nsew", padx=8, pady=8)

        tk.Label(f_ai, text="AI 분석 결과를 붙여넣으세요.").pack(pady=4)

        r_tag = tk.Frame(f_ai)
        r_tag.pack(fill="x", padx=10, pady=2)
        tk.Label(r_tag, text="취약점 키워드:").pack(side="left")
        self.ent_weakness = tk.Entry(r_tag, width=15, font=("맑은 고딕", 10, "bold"), fg="#dc2626")
        self.ent_weakness.pack(side="left", padx=4)

        r_score = tk.Frame(f_ai)
        r_score.pack(fill="x", padx=10, pady=2)
        tk.Label(r_score, text="이해도 (0~100):").pack(side="left")
        self.ent_score = tk.Entry(r_score, width=6)
        self.ent_score.pack(side="left", padx=4)

        self.txt_feedback = tk.Text(f_ai, height=12, bg="#f0f9ff", font=("맑은 고딕", 10))
        self.txt_feedback.pack(fill="both", expand=True, padx=10, pady=4)

        tk.Button(
            f_ai, text="분석 결과 저장", bg="#ea580c", fg="white",
            font=("맑은 고딕", 11, "bold"), height=2,
            command=self._save_analysis,
        ).pack(fill="x", padx=16, pady=8)

        # 우: 시각화
        f_viz = tk.LabelFrame(self.tab_analysis, text=" 3. 시각화 & 처방 ", font=("맑은 고딕", 10, "bold"))
        f_viz.grid(row=0, column=2, sticky="nsew", padx=8, pady=8)

        tk.Button(
            f_viz, text="성취도 추이 (Line)", bg="#0891b2", fg="white",
            font=("맑은 고딕", 10, "bold"), height=2,
            command=self._chart_progress,
        ).pack(fill="x", padx=16, pady=4)
        tk.Button(
            f_viz, text="취약점 분석 (Bar)", bg="#7c3aed", fg="white",
            font=("맑은 고딕", 10, "bold"), height=2,
            command=self._chart_weakness,
        ).pack(fill="x", padx=16, pady=4)
        tk.Button(
            f_viz, text="성장 리포트 CSV 내보내기", bg="#334155", fg="white",
            font=("맑은 고딕", 10, "bold"), height=2,
            command=self._export_csv,
        ).pack(fill="x", padx=16, pady=12)

        tk.Label(f_viz, text="타겟 취약점:").pack(pady=(8, 0))
        self.lbl_weakness = tk.Label(f_viz, text="없음", font=("Consolas", 12, "bold"), fg="#dc2626")
        self.lbl_weakness.pack(pady=4)

        tk.Button(
            f_viz, text="취약점 타겟 보충문제 자동 추출", bg="#16a34a", fg="white",
            font=("맑은 고딕", 11, "bold"), height=3,
            command=self._make_remedial,
        ).pack(fill="x", padx=16, pady=4)

    # ── 학생 CRUD ─────────────────────────────
    def _add_student(self):
        name = self.ent_student.get().strip()
        if not name:
            return messagebox.showwarning("오류", "이름을 입력하세요.")
        with db_conn(self.cfg) as conn:
            conn.cursor().execute("INSERT INTO students (name, grade) VALUES (?,?)", (name, "공통"))
        self.ent_student.delete(0, tk.END)
        self._refresh_students()
        self._refresh_status()

    def _delete_student(self):
        if not self.current_student_id:
            return messagebox.showwarning("알림", "학생을 선택하세요.")
        if not messagebox.askyesno("확인", f"'{self.current_student_name}' 학생과 모든 분석 기록을 삭제하시겠습니까?"):
            return
        with db_conn(self.cfg) as conn:
            cur = conn.cursor()
            cur.execute("DELETE FROM analysis_records WHERE student_id=?", (self.current_student_id,))
            cur.execute("DELETE FROM students WHERE id=?", (self.current_student_id,))
        self.current_student_id = None
        self.current_student_name = ""
        self._refresh_students()
        self._refresh_status()
        self.txt_feedback.delete("1.0", tk.END)
        self.lbl_weakness.config(text="없음")

    def _rename_student(self):
        if not self.current_student_id:
            return messagebox.showwarning("알림", "학생을 선택하세요.")
        from tkinter import simpledialog
        new_name = simpledialog.askstring("이름 수정", "새 이름:", parent=self.root, initialvalue=self.current_student_name)
        if not new_name or not new_name.strip():
            return
        with db_conn(self.cfg) as conn:
            conn.cursor().execute("UPDATE students SET name=? WHERE id=?", (new_name.strip(), self.current_student_id))
        self.current_student_name = new_name.strip()
        self._refresh_students()

    def _refresh_students(self):
        self.list_students.delete(0, tk.END)
        with db_conn(self.cfg) as conn:
            for row in conn.cursor().execute("SELECT id,name FROM students ORDER BY id"):
                self.list_students.insert(tk.END, f"[{row[0]}] {row[1]}")

    def _on_student_select(self, event):
        sel = self.list_students.curselection()
        if not sel:
            return
        val = self.list_students.get(sel[0])
        self.current_student_id = int(re.search(r"\[(\d+)\]", val).group(1))
        self.current_student_name = re.search(r"\]\s*(.*)", val).group(1)

        self.txt_feedback.delete("1.0", tk.END)
        self.txt_feedback.insert(tk.END, f"--- [{self.current_student_name}] 누적 히스토리 ---\n\n")

        with db_conn(self.cfg) as conn:
            history = conn.cursor().execute(
                "SELECT record_date,weakness_tag,score,feedback FROM analysis_records WHERE student_id=? ORDER BY id DESC",
                (self.current_student_id,),
            ).fetchall()

        if history:
            self.analyzed_weakness = history[0][1]
            self.lbl_weakness.config(text=f"Target: {self.analyzed_weakness}")
            for h in history:
                self.txt_feedback.insert(tk.END, f"Date: {h[0]} | Tag: {h[1]} | Score: {h[2]}\n{h[3]}\n\n")
        else:
            self.analyzed_weakness = None
            self.lbl_weakness.config(text="없음")

    def _copy_ai_prompt(self):
        if not self.current_student_id:
            return messagebox.showwarning("알림", "학생을 선택하세요.")
        prompt = (
            f"[STUDENT_DATA_INPUT]\n"
            f"Student_ID: {self.current_student_id}\n"
            f"Student_Name: {self.current_student_name}\n"
            f"Test_Date: {datetime.date.today().isoformat()}\n"
            f"[CURRENT_TEST_DATA]\n"
            f"(여기에 학생의 오답 내용을 붙여넣어 주세요)\n"
            f"[/STUDENT_DATA_INPUT]"
        )
        self.root.clipboard_clear()
        self.root.clipboard_append(prompt)
        messagebox.showinfo("복사 완료", f"'{self.current_student_name}' 프롬프트 복사 완료!")

    def _save_analysis(self):
        if not self.current_student_id:
            return messagebox.showwarning("오류", "학생을 선택하세요.")
        weakness = self.ent_weakness.get().strip()
        score_str = self.ent_score.get().strip()
        feedback = self.txt_feedback.get("1.0", tk.END).strip()
        if not weakness or not score_str.isdigit():
            return messagebox.showwarning("오류", "취약점 키워드와 점수(숫자)를 입력하세요.")
        with db_conn(self.cfg) as conn:
            conn.cursor().execute(
                "INSERT INTO analysis_records (student_id,record_date,weakness_tag,score,feedback) VALUES (?,?,?,?,?)",
                (self.current_student_id, datetime.date.today().isoformat(), weakness, int(score_str), feedback),
            )
        self.ent_weakness.delete(0, tk.END)
        self.ent_score.delete(0, tk.END)
        messagebox.showinfo("저장 완료", "누적 저장 완료!")
        self._on_student_select(None)

    # ── 차트 ──────────────────────────────────
    def _chart_progress(self):
        if not self.current_student_id:
            return messagebox.showwarning("오류", "학생을 선택하세요.")
        with db_conn(self.cfg) as conn:
            recs = conn.cursor().execute(
                "SELECT record_date,score FROM analysis_records WHERE student_id=? ORDER BY id",
                (self.current_student_id,),
            ).fetchall()
        if not recs:
            return messagebox.showinfo("안내", "데이터가 없습니다.")

        dates = [f"{r[0][5:]}({i+1}차)" for i, r in enumerate(recs)]
        scores = [r[1] for r in recs]

        fig, ax = plt.subplots(figsize=(9, 5))
        ax.plot(dates, scores, marker="o", color="#2563eb", linewidth=2.5, markersize=8)
        ax.fill_between(dates, scores, alpha=0.12, color="#2563eb")
        ax.axhline(y=80, color="#dc2626", linestyle="--", alpha=0.6, label="목표 80점")
        for i, s in enumerate(scores):
            ax.annotate(f"{s}", (dates[i], s), textcoords="offset points", xytext=(0, 10), ha="center", fontsize=9, fontweight="bold")
        ax.set_title(f"[{self.current_student_name}] 성취도 변화", fontsize=15, fontweight="bold", pad=15)
        ax.set_xlabel("평가 시기", fontsize=11)
        ax.set_ylabel("이해도 (100점)", fontsize=11)
        ax.set_ylim(0, 105)
        ax.legend()
        fig.tight_layout()
        plt.show()

    def _chart_weakness(self):
        if not self.current_student_id:
            return messagebox.showwarning("오류", "학생을 선택하세요.")
        with db_conn(self.cfg) as conn:
            recs = conn.cursor().execute(
                "SELECT weakness_tag FROM analysis_records WHERE student_id=?",
                (self.current_student_id,),
            ).fetchall()
        if not recs:
            return messagebox.showinfo("안내", "데이터가 없습니다.")

        counts = Counter(r[0] for r in recs if r[0] and r[0].strip())
        items = counts.most_common()
        labels = [x[0] for x in items][::-1]
        values = [x[1] for x in items][::-1]

        fig, ax = plt.subplots(figsize=(9, 5))
        bars = ax.barh(labels, values, color="#7c3aed", alpha=0.85, height=0.6)
        for b in bars:
            ax.text(b.get_width() + 0.15, b.get_y() + b.get_height() / 2,
                    f"{int(b.get_width())}건", va="center", fontsize=11, fontweight="bold")
        ax.set_title(f"[{self.current_student_name}] 취약점 분석", fontsize=15, fontweight="bold", pad=15)
        ax.set_xlabel("횟수", fontsize=11)
        if values:
            ax.set_xticks(range(0, max(values) + 2))
        fig.tight_layout()
        plt.show()

    def _export_csv(self):
        if not self.current_student_id:
            return messagebox.showwarning("오류", "학생을 선택하세요.")
        path = os.path.join(self.cfg["last_dir"], f"[{self.current_student_name}]_학습리포트.csv")
        with db_conn(self.cfg) as conn:
            recs = conn.cursor().execute(
                "SELECT record_date,weakness_tag,score,feedback FROM analysis_records WHERE student_id=? ORDER BY id",
                (self.current_student_id,),
            ).fetchall()
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["날짜", "학생", "취약점", "이해도", "분석 내용"])
            for r in recs:
                w.writerow([r[0], self.current_student_name, r[1], f"{r[2]}점", str(r[3]).strip()])
        messagebox.showinfo("완료", f"CSV 저장: {path}")
        open_directory(self.cfg["last_dir"])

    def _make_remedial(self):
        if not self.analyzed_weakness:
            return messagebox.showwarning("경고", "먼저 AI 분석 결과를 저장하세요.")
        kw = self.analyzed_weakness.replace("구문", "").strip()

        with db_conn(self.cfg) as conn:
            cur = conn.cursor()
            matches = cur.execute("SELECT id FROM passages WHERE extra_tags LIKE ?", (f"%{kw}%",)).fetchall()

        if not matches:
            return messagebox.showinfo("결과", f"'{kw}' 관련 문제가 DB에 없습니다.")

        self._cart_clear()
        for m in matches:
            self.cart[m[0]] = set()
        self._sync_cart_ui()
        messagebox.showinfo("추출 완료", f"{len(matches)}개 지문 추출.\n문제 관리 탭에서 생성하세요.")
        self.tabs.select(self.tab_bank)

    # ──────────────────────────────────────────────
    #  TAB 4: 설정 및 백업
    # ──────────────────────────────────────────────
    def _build_config_tab(self):
        main = tk.Frame(self.tab_config)
        main.pack(fill="both", expand=True, padx=30, pady=20)

        # 경로 설정
        path_frame = tk.LabelFrame(main, text=" 경로 설정 ", font=("맑은 고딕", 10, "bold"))
        path_frame.pack(fill="x", pady=(0, 12))

        r0 = tk.Frame(path_frame)
        r0.pack(fill="x", padx=10, pady=6)
        tk.Label(r0, text="출력 폴더:", width=18, anchor="w").pack(side="left")
        self.lbl_out = tk.Label(r0, text=self.cfg["last_dir"], fg="#2563eb")
        self.lbl_out.pack(side="left", padx=8)
        tk.Button(r0, text="변경", command=self._change_out_dir).pack(side="right", padx=4)
        tk.Button(r0, text="열기", command=lambda: open_directory(self.cfg["last_dir"])).pack(side="right", padx=4)

        r1 = tk.Frame(path_frame)
        r1.pack(fill="x", padx=10, pady=6)
        tk.Label(r1, text="DB 폴더:", width=18, anchor="w").pack(side="left")
        self.lbl_db = tk.Label(r1, text=self.cfg.get("db_dir", SCRIPT_DIR), fg="#16a34a")
        self.lbl_db.pack(side="left", padx=8)
        tk.Button(r1, text="변경", command=self._change_db_dir).pack(side="right", padx=4)
        tk.Button(r1, text="열기", command=lambda: open_directory(self.cfg.get("db_dir", SCRIPT_DIR))).pack(side="right", padx=4)

        # 백업/복원
        backup_frame = tk.LabelFrame(main, text=" DB 백업 및 복원 ", font=("맑은 고딕", 10, "bold"))
        backup_frame.pack(fill="x", pady=12)

        br = tk.Frame(backup_frame)
        br.pack(fill="x", padx=10, pady=10)
        tk.Button(
            br, text="DB 백업 생성", bg="#2563eb", fg="white",
            font=("맑은 고딕", 10, "bold"), width=18,
            command=self._backup_db,
        ).pack(side="left", padx=8)
        tk.Button(
            br, text="백업에서 복원", bg="#ea580c", fg="white",
            font=("맑은 고딕", 10, "bold"), width=18,
            command=self._restore_db,
        ).pack(side="left", padx=8)

        self.lbl_backup_info = tk.Label(backup_frame, text="", font=("맑은 고딕", 9), fg="#64748b")
        self.lbl_backup_info.pack(padx=10, pady=(0, 8))

        # 최근 백업 표시
        self._update_backup_info()

        # DB 통계
        stat_frame = tk.LabelFrame(main, text=" DB 현황 ", font=("맑은 고딕", 10, "bold"))
        stat_frame.pack(fill="x", pady=12)
        try:
            s = db_stats(self.cfg)
            stat_text = f"지문: {s['passages']}개   문제: {s['questions']}개   학생: {s['students']}명"
        except Exception:
            stat_text = "DB 연결 오류"
        tk.Label(stat_frame, text=stat_text, font=("맑은 고딕", 11), pady=10).pack()

    def _change_out_dir(self):
        d = filedialog.askdirectory()
        if d:
            self.cfg["last_dir"] = d
            self.lbl_out.config(text=d)
            save_config(self.cfg)

    def _change_db_dir(self):
        d = filedialog.askdirectory()
        if d:
            self.cfg["db_dir"] = d
            self.lbl_db.config(text=d)
            save_config(self.cfg)
            init_db(self.cfg)
            self._refresh_bank()
            self._refresh_status()
            messagebox.showinfo("완료", "DB 경로 변경 완료.")

    def _backup_db(self):
        src = get_db_path(self.cfg)
        if not os.path.exists(src):
            return messagebox.showwarning("오류", "DB 파일이 없습니다.")
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        dst = src.replace(".db", f"_backup_{ts}.db")
        try:
            shutil.copy2(src, dst)
            messagebox.showinfo("백업 완료", f"백업 파일:\n{dst}")
            self._update_backup_info()
        except Exception as e:
            messagebox.showerror("오류", f"백업 실패: {e}")

    def _restore_db(self):
        src = filedialog.askopenfilename(
            title="복원할 백업 파일 선택",
            filetypes=[("SQLite DB", "*.db")],
            initialdir=self.cfg.get("db_dir", SCRIPT_DIR),
        )
        if not src:
            return
        if not messagebox.askyesno("경고", "현재 DB를 선택한 백업으로 덮어씁니다.\n계속하시겠습니까?"):
            return
        dst = get_db_path(self.cfg)
        try:
            shutil.copy2(src, dst)
            init_db(self.cfg)
            self._refresh_bank()
            self._refresh_students()
            self._refresh_status()
            messagebox.showinfo("복원 완료", "DB 복원이 완료되었습니다.")
        except Exception as e:
            messagebox.showerror("오류", f"복원 실패: {e}")

    def _update_backup_info(self):
        db_dir = self.cfg.get("db_dir", SCRIPT_DIR)
        backups = sorted(
            [f for f in os.listdir(db_dir) if f.endswith(".db") and "backup" in f],
            reverse=True,
        ) if os.path.isdir(db_dir) else []
        if backups:
            self.lbl_backup_info.config(text=f"최근 백업: {backups[0]}")
        else:
            self.lbl_backup_info.config(text="백업 파일 없음")


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 엔트리포인트
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
if __name__ == "__main__":
    os.chdir(SCRIPT_DIR)
    root = tk.Tk()
    app = ExamProApp(root)
    root.mainloop()
