"""
문서 생성 엔진 — Word/PDF 시험지, 워크북, 정답지
"""
import os, re, time, datetime, platform, json
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
try:
    from docx2pdf import convert as _convert_pdf
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

def _setup_doc(doc, cfg, title):
    style = doc.styles["Normal"]
    style.font.name = cfg["font_name"]
    style._element.rPr.rFonts.set(qn("w:eastAsia"), cfg["font_name"])
    style.font.size = Pt(int(cfg["font_size"]))
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    sec.left_margin = sec.right_margin = Cm(1.5)
    doc.add_heading(title, 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

def convert_to_pdf_safe(docx_path, pdf_path):
    time.sleep(1.0)
    if platform.system() == "Windows":
        try:
            import win32com.client
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(docx_path, ReadOnly=True)
            doc.SaveAs(pdf_path, FileFormat=17)
            doc.Close(); word.Quit()
            return True, "성공"
        except Exception as e:
            return False, f"Word COM 오류: {e}"
    else:
        if not PDF_AVAILABLE: return False, "docx2pdf 미설치"
        try: _convert_pdf(docx_path, pdf_path); return True, "성공"
        except Exception as e: return False, str(e)

def create_exam_files(target_dir, prefix, exam_data, cfg,
                      is_workbook=False, output_format="Word + PDF",
                      student_name="", set_label=""):
    ts = datetime.datetime.now().strftime("%m%d_%H%M")
    label = f"_{set_label}" if set_label else ""
    doc_q = docx.Document()
    title_prefix = "워크북" if is_workbook else "시험지"
    title_text = f"실전 대비 맞춤형 {title_prefix}"
    if student_name: title_text += f" [{student_name}]"
    _setup_doc(doc_q, cfg, title_text)
    if not is_workbook:
        new_sec = doc_q.add_section(WD_SECTION.CONTINUOUS)
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), "2"); cols.set(qn("w:space"), "708")
        new_sec._sectPr.append(cols)
    g_idx = 1
    for idx, item in enumerate(exam_data, 1):
        if not item["passage"]: continue
        item["q_mappings"] = []
        if is_workbook:
            h = doc_q.add_paragraph()
            h.add_run(f"■ 지문 {idx}  ").bold = True
            h.add_run(f"[{item.get('info', '')}]").font.color.rgb = RGBColor(100, 100, 100)
            sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", item["passage"].strip()) if s.strip()]
            for s_i, sent in enumerate(sentences, 1):
                p = doc_q.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.add_run(f"{s_i}. {sent}").bold = True
                for _ in range(2):
                    bl = doc_q.add_paragraph()
                    bl.paragraph_format.space_after = Pt(0)
                    bl.add_run("_" * 74).font.color.rgb = RGBColor(220, 220, 220)
            doc_q.add_paragraph("\n")
        else:
            intro = doc_q.add_paragraph()
            intro.add_run("※ 다음 글을 읽고 물음에 답하시오.").bold = True
            tbl = doc_q.add_table(rows=1, cols=1)
            tbl.style = "Table Grid"
            tbl.cell(0, 0).text = item["passage"]
            doc_q.add_paragraph()
            for q in item["questions"]:
                p = doc_q.add_paragraph()
                p.add_run(f"{g_idx}. ").bold = True
                p.add_run(q["text"])
                choices = q.get("choices", [])
                if choices and any(c for c in choices if c):
                    if isinstance(choices, str):
                        try: choices = json.loads(choices)
                        except json.JSONDecodeError: choices = []
                    for ch in choices:
                        if ch: doc_q.add_paragraph(f"    {ch}")
                doc_q.add_paragraph()
                item["q_mappings"].append({
                    "new_num": g_idx,
                    "orig_num": q.get("num", q.get("q_num", "-")),
                    "preview": q["text"][:25].replace("\n", " ") + "...",
                })
                g_idx += 1
    q_docx = os.path.normpath(os.path.join(target_dir, f"{prefix}_문제지_{ts}{label}.docx"))
    doc_q.save(q_docx)
    doc_a = docx.Document()
    _setup_doc(doc_a, cfg, f"맞춤형 {title_prefix} — 정답 및 해설")
    for idx, item in enumerate(exam_data, 1):
        h = doc_a.add_paragraph()
        h.add_run(f"■ [지문 {idx}] 해설\n").bold = True
        if not is_workbook and item.get("q_mappings"):
            pm = doc_a.add_paragraph()
            rt = pm.add_run("[※ 문항 번호 매칭표]\n"); rt.bold = True
            rt.font.color.rgb = RGBColor(0, 112, 192)
            lines = ""
            for m in item["q_mappings"]:
                orig = f"{m['orig_num']}번" if m["orig_num"] != "-" else "서술형"
                lines += f"▶ 시험지 {m['new_num']}번 = 본문 해설 {orig} 참조 | {m['preview']}\n"
            rm = pm.add_run(lines)
            rm.font.color.rgb = RGBColor(100, 100, 100); rm.font.size = Pt(9)
            doc_a.add_paragraph()
        body = doc_a.add_paragraph()
        ans = (item.get("answer_text") or "").strip()
        body.add_run(ans if ans else "(입력된 정답/해설이 없습니다.)")
        doc_a.add_paragraph("\n" + "=" * 50 + "\n")
    a_docx = os.path.normpath(os.path.join(target_dir, f"{prefix}_정답지_{ts}{label}.docx"))
    doc_a.save(a_docx)
    msg = "Word 파일(문제/답안) 분리 생성 완료!"
    if output_format in ("PDF만", "Word + PDF"):
        q_pdf = q_docx.replace(".docx", ".pdf"); a_pdf = a_docx.replace(".docx", ".pdf")
        q_ok, q_err = convert_to_pdf_safe(q_docx, q_pdf)
        a_ok, a_err = convert_to_pdf_safe(a_docx, a_pdf)
        if q_ok and a_ok:
            msg = "Word + PDF 분리 생성 완료!"
            if output_format == "PDF만":
                for f in (q_docx, a_docx):
                    try: os.remove(f)
                    except OSError: pass
                msg = "PDF 분리 생성 완료!"
        else: msg = f"Word 생성 완료. PDF 변환 실패: {q_err}"
    return msg
