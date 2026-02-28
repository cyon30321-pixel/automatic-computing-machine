"""
공통 Word 문서 유틸리티 v6.0
— generator.py와 vocab_generator.py가 공유하는 함수들
— 파일명 sanitization, 배너, 스타일, PDF 변환, 셀 서식
"""

import os
import re
import time
import hashlib
import platform

import docx
from docx.shared import Cm, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    from docx2pdf import convert as _convert_pdf
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

from exam_bank.constants import Theme

MAX_SAFE_PATH = 240

# Windows 금지 문자 패턴
_INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')


# ── 파일 경로 유틸리티 ──

def sanitize_filename(name):
    """파일명에서 Windows 금지 문자 제거"""
    if not name:
        return name
    sanitized = _INVALID_FILENAME_CHARS.sub("_", name)
    sanitized = sanitized.strip(". ")
    if not sanitized:
        sanitized = "untitled"
    return sanitized


def safe_basename(folder, stem, ext):
    full = os.path.join(folder, stem + ext)
    if len(full) <= MAX_SAFE_PATH:
        return stem
    h = hashlib.md5(stem.encode("utf-8")).hexdigest()[:6]
    cut_len = max(1, MAX_SAFE_PATH - (len(folder) + 1 + len(ext) + 8))
    return f"{stem[:cut_len]}_{h}"


def build_paths(target_dir, base_filename, prefix=""):
    stem_raw = f"{prefix}{sanitize_filename(base_filename)}"
    safe_stem = safe_basename(target_dir, stem_raw, ".docx")
    docx_path = os.path.join(target_dir, safe_stem + ".docx")
    pdf_path = os.path.join(target_dir, safe_stem + ".pdf")
    return docx_path, pdf_path


def safe_save_docx(doc, docx_path):
    """방어 — 파일이 열려있으면 명확한 에러 반환"""
    try:
        doc.save(docx_path)
        return True, ""
    except PermissionError:
        return False, f"파일이 다른 프로그램에서 열려 있습니다.\n먼저 닫아주세요:\n{os.path.basename(docx_path)}"
    except Exception as e:
        return False, str(e)


# ── PDF 변환 ──

def convert_to_pdf_safe(docx_path, pdf_path):
    time.sleep(0.5)
    if platform.system() == "Windows":
        try:
            import win32com.client
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            d = word.Documents.Open(os.path.abspath(docx_path), ReadOnly=True)
            d.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
            d.Close()
            word.Quit()
            return True, "success"
        except Exception as e:
            if PDF_AVAILABLE:
                try:
                    _convert_pdf(docx_path, pdf_path)
                    return True, "success"
                except Exception:
                    pass
            return False, f"PDF conversion error: {e}"
    else:
        if not PDF_AVAILABLE:
            return False, "docx2pdf not installed"
        try:
            _convert_pdf(docx_path, pdf_path)
            return True, "success"
        except Exception as e:
            return False, str(e)


def convert_docx_pairs_to_pdf_checked(pairs):
    results = []
    for docx_path, pdf_path in pairs:
        if not os.path.exists(docx_path):
            continue
        ok, msg = convert_to_pdf_safe(docx_path, pdf_path)
        results.append((ok, msg))
    return results


# ── 문서 스타일 ──

def apply_doc_style(doc, font_name, font_size, margins=None):
    """기본 문서 스타일 적용 (폰트, 여백)"""
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(int(font_size))
    style.font.color.rgb = RGBColor.from_string(Theme.DARK_TEXT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    section = doc.sections[0]
    if margins:
        section.top_margin = Cm(margins.get("top", 1.5))
        section.bottom_margin = Cm(margins.get("bottom", 1.5))
        section.left_margin = Cm(margins.get("left", 1.5))
        section.right_margin = Cm(margins.get("right", 1.5))
    else:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)


def set_two_columns(section):
    try:
        sectPr = section._sectPr
        cols = sectPr.xpath("./w:cols")
        cols_elem = cols[0] if cols else OxmlElement("w:cols")
        cols_elem.set(qn("w:num"), "2")
        cols_elem.set(qn("w:space"), "600")
        cols_elem.set(qn("w:sep"), "1")
        if not cols:
            sectPr.append(cols_elem)
    except Exception:
        pass


# ── 셀/테이블 서식 ──

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)


def set_cell_borders(cell, color="B0C4DE", size="4"):
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), size)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tcBorders.append(border)
    tcPr.append(tcBorders)


def set_cell_vertical_alignment(cell, align="center"):
    tcPr = cell._element.get_or_add_tcPr()
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), align)
    tcPr.append(vAlign)


def set_cell_margins(cell, top='100', bottom='100', left='140', right='140'):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        elem = OxmlElement(f'w:{side}')
        elem.set(qn('w:w'), val)
        elem.set(qn('w:type'), 'dxa')
        tcMar.append(elem)
    tcPr.append(tcMar)


def set_paragraph_border_bottom(paragraph, color="2E86C1", size="6", space="3"):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), space)
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


# ── 텍스트 서식 ──

def add_formatted_text(paragraph, text, font_size=None, color=None):
    pattern = r'([①-⑳ⓐ-ⓩⒶ-Ⓩ]|\([A-Ea-e]\))'
    tokens = re.split(pattern, text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run(token)
        if font_size:
            run.font.size = Pt(font_size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if re.match(pattern, token):
            run.bold = True


# ── 배너/헤더 ──

def add_header_banner(doc, title_text, subtitle_text="", logo_path="", font_name="맑은 고딕"):
    if logo_path and os.path.exists(logo_path):
        logo_p = doc.add_paragraph()
        logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            logo_p.add_run().add_picture(logo_path, height=Cm(2.5))
        except Exception:
            pass
        logo_p.paragraph_format.space_after = Pt(4)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(6)
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run(title_text)
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor.from_string(Theme.PRIMARY)
    title_run.font.name = font_name
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)

    if subtitle_text:
        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sub_p.add_run(subtitle_text)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(Theme.GRAY_TEXT)

    sep_p = doc.add_paragraph()
    sep_p.paragraph_format.space_before = Pt(0)
    sep_p.paragraph_format.space_after = Pt(8)
    set_paragraph_border_bottom(sep_p, color=Theme.SECONDARY, size="8")
    return title_p


def set_cell_width(cell, width_cm):
    """셀 너비를 cm 단위로 설정"""
    tcPr = cell._element.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(int(width_cm * 567)))  # cm -> twips
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)
