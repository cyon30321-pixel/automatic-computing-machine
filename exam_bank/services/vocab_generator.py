"""
단어 시험지 생성 엔진 v6.0
— FactoryVoca Pro 스타일 2단 레이아웃
— 영→한 / 한→영 모드, 정답지, 첨글자 힌트
"""

import os
import random
import datetime

import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from exam_bank.constants import Theme
from exam_bank.services.docx_utils import (
    apply_doc_style, add_header_banner, build_paths,
    safe_save_docx, convert_to_pdf_safe, sanitize_filename,
    set_cell_borders, set_cell_background, set_cell_vertical_alignment,
    set_cell_margins, set_paragraph_border_bottom, set_cell_width,
)


def _make_answer_hint(word, show_first_letter=False):
    """첨글자 힌트 생성"""
    if not show_first_letter or not word:
        return ""
    return word[0] + "_" * (len(word) - 1)


def _build_vocab_table(doc, words, exam_type, show_answer=False,
                       show_first_letter=False, font_name="맑은 고딕"):
    """
    FactoryVoca 스타일 2단 단어 테이블 생성
    words: [{"english": ..., "korean": ..., "pos": ...}, ...]
    """
    total = len(words)
    if total == 0:
        doc.add_paragraph("(출제할 단어가 없습니다)")
        return

    # 2단 분할: 왼쪽 1~half, 오른쪽 half+1~total
    half = (total + 1) // 2
    left_words = words[:half]
    right_words = words[half:]

    # 테이블: [번호 | 출제어 | 답란] [번호 | 출제어 | 답란]
    # 6열 구조
    rows_needed = max(len(left_words), len(right_words))
    table = doc.add_table(rows=rows_needed + 1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 헤더 행
    header_labels = ["No.", "단어", "뜻"] if exam_type == "eng_to_kor" else ["No.", "뜻", "단어"]
    header_labels_full = header_labels + header_labels
    header_row = table.rows[0]
    for i, label in enumerate(header_labels_full):
        cell = header_row.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, Theme.PRIMARY)
        set_cell_borders(cell, color=Theme.TABLE_BORDER, size="4")
        set_cell_vertical_alignment(cell)

    # 데이터 행
    for row_idx in range(rows_needed):
        row = table.rows[row_idx + 1]
        for side in range(2):  # 0=왼쪽, 1=오른쪽
            word_list = left_words if side == 0 else right_words
            col_offset = side * 3
            if row_idx < len(word_list):
                w = word_list[row_idx]
                num = row_idx + 1 if side == 0 else half + row_idx + 1

                if exam_type == "eng_to_kor":
                    question_text = w.get("english", "")
                    pos_text = f" ({w.get('pos', '')})" if w.get("pos") else ""
                    question_text += pos_text
                    answer_text = w.get("korean", "")
                else:
                    question_text = w.get("korean", "")
                    answer_text = w.get("english", "")

                # 번호 셀
                num_cell = row.cells[col_offset]
                num_cell.text = ""
                p = num_cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(str(num))
                run.font.size = Pt(9)
                run.font.name = font_name
                run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
                set_cell_borders(num_cell, color=Theme.TABLE_BORDER, size="2")
                set_cell_vertical_alignment(num_cell)

                # 출제어 셀
                q_cell = row.cells[col_offset + 1]
                q_cell.text = ""
                p = q_cell.paragraphs[0]
                run = p.add_run(question_text)
                run.font.size = Pt(9)
                run.font.name = font_name
                run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
                set_cell_borders(q_cell, color=Theme.TABLE_BORDER, size="2")
                set_cell_vertical_alignment(q_cell)
                set_cell_margins(q_cell, left='80', right='80')

                # 답란 셀
                a_cell = row.cells[col_offset + 2]
                a_cell.text = ""
                p = a_cell.paragraphs[0]
                if show_answer:
                    run = p.add_run(answer_text)
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor.from_string(Theme.ACCENT)
                    run.font.name = font_name
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
                elif show_first_letter:
                    hint = _make_answer_hint(answer_text, True)
                    run = p.add_run(hint)
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(180, 180, 180)
                    run.font.name = font_name
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
                set_cell_borders(a_cell, color=Theme.TABLE_BORDER, size="2")
                set_cell_vertical_alignment(a_cell)
                set_cell_margins(a_cell, left='80', right='80')
            else:
                # 빈 셀
                for c in range(3):
                    cell = row.cells[col_offset + c]
                    cell.text = ""
                    set_cell_borders(cell, color=Theme.TABLE_BORDER, size="2")

    # 열 너비 설정
    for row in table.rows:
        for side in range(2):
            offset = side * 3
            set_cell_width(row.cells[offset], 1.0)       # 번호
            set_cell_width(row.cells[offset + 1], 3.8)   # 출제어
            set_cell_width(row.cells[offset + 2], 3.8)   # 답란

    return table


def create_vocab_test(target_dir, words, cfg,
                      exam_type="eng_to_kor",
                      student_name="",
                      show_answer=False,
                      show_first_letter=False,
                      shuffle=False,
                      custom_title="",
                      custom_filename="",
                      logo_path="",
                      output_format="Word + PDF",
                      unit_names=None):
    """
    단어 시험지 생성 메인 함수

    words: [{"english": ..., "korean": ..., "pos": ...}, ...]
    unit_names: 출제 단원 이름 목록 (부제에 표시)
    반환: (생성 파일 경로 리스트, 에러 메시지)
    """
    if not words:
        return [], "출제할 단어가 없습니다."

    os.makedirs(target_dir, exist_ok=True)

    font_name = cfg.get("font_name", "맑은 고딕")
    font_size = cfg.get("font_size", "10")
    if not logo_path:
        logo_path = cfg.get("last_logo", "")

    ts = datetime.datetime.now().strftime("%m%d_%H%M")
    mode_label = "영한" if exam_type == "eng_to_kor" else "한영"

    # 제목 결정
    if custom_title:
        title_text = custom_title
    else:
        title_text = "Vocabulary Test"
        if student_name:
            title_text += f" [{student_name}]"

    # 파일명 결정
    if custom_filename:
        base_name = sanitize_filename(custom_filename)
    else:
        name_part = f"_{student_name}" if student_name else ""
        base_name = f"VocabTest_{mode_label}{name_part}_{ts}"

    # 부제: 출제 범위
    subtitle_parts = []
    if unit_names:
        if len(unit_names) <= 3:
            subtitle_parts.append(" / ".join(unit_names))
        else:
            subtitle_parts.append(f"{unit_names[0]} ~ {unit_names[-1]} ({len(unit_names)}개 단원)")
    subtitle_parts.append(f"{mode_label} | {len(words)}문항")
    subtitle_text = "  |  ".join(subtitle_parts)

    # 단어 셔플
    if shuffle:
        words = list(words)
        random.shuffle(words)

    created_files = []
    errors = []

    # ── 시험지 생성 ──
    doc_q = docx.Document()
    apply_doc_style(doc_q, font_name, font_size)
    add_header_banner(doc_q, title_text, subtitle_text, logo_path, font_name)

    # 학생 정보 란
    info_p = doc_q.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    info_p.paragraph_format.space_before = Pt(4)
    info_p.paragraph_format.space_after = Pt(8)
    date_str = datetime.datetime.now().strftime("%Y. %m. %d")
    run = info_p.add_run(f"Name: ________________    Score: ______ / {len(words)}    Date: {date_str}")
    run.font.size = Pt(10)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)

    _build_vocab_table(doc_q, words, exam_type,
                       show_answer=False,
                       show_first_letter=show_first_letter,
                       font_name=font_name)

    # 시험지 저장
    q_docx, q_pdf = build_paths(target_dir, base_name)
    ok, err = safe_save_docx(doc_q, q_docx)
    if not ok:
        return [], err
    created_files.append(q_docx)

    # ── 정답지 생성 ──
    ans_base = f"{base_name}_정답"
    doc_a = docx.Document()
    apply_doc_style(doc_a, font_name, font_size)
    add_header_banner(doc_a, f"{title_text} [정답지]", subtitle_text, logo_path, font_name)

    _build_vocab_table(doc_a, words, exam_type,
                       show_answer=True,
                       font_name=font_name)

    a_docx, a_pdf = build_paths(target_dir, ans_base)
    ok, err = safe_save_docx(doc_a, a_docx)
    if not ok:
        errors.append(f"정답지 저장 실패: {err}")
    else:
        created_files.append(a_docx)

    # ── PDF 변환 ──
    if "PDF" in output_format:
        for dx, px in [(q_docx, q_pdf), (a_docx, a_pdf)]:
            if os.path.exists(dx):
                ok, msg = convert_to_pdf_safe(dx, px)
                if ok:
                    created_files.append(px)
                else:
                    errors.append(f"PDF 변환 실패: {msg}")

    err_msg = "\n".join(errors) if errors else ""
    return created_files, err_msg
